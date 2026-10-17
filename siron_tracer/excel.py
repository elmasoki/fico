"""Excel round trip for the mapping: export a styled workbook that business
analysts can fill in (status, target field, note, sensitive), and import the
edited workbook back."""
from __future__ import annotations

import io
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

STATUS_LABELS = {"migrate": "Taşınacak", "review": "İncelenecek", "ignore": "Gereksiz"}
STATUS_CODES = {v.casefold(): k for k, v in STATUS_LABELS.items()}
STATUS_CODES.update({k: k for k in STATUS_LABELS})
YES = {"evet", "e", "yes", "y", "1", "true", "x"}
NO = {"hayır", "hayir", "h", "no", "n", "0", "false"}

HEADER_FILL = PatternFill("solid", fgColor="14213D")
EDIT_FILL = PatternFill("solid", fgColor="FFF8E1")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D0D5DD")


def _sheet(wb: Workbook, title: str, headers: List[str], rows: List[list], widths: Dict[int, int] = None,
           editable: List[str] = (), first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    edit_idx = [headers.index(h) + 1 for h in editable if h in headers]
    for r in range(2, ws.max_row + 1):
        for c in edit_idx:
            ws.cell(r, c).fill = EDIT_FILL
        for c in range(1, len(headers) + 1):
            ws.cell(r, c).border = Border(bottom=THIN)
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=c in edit_idx)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(12, min(40, len(h) + 4)))
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), ws.max_row)
    ws.row_dimensions[1].height = 30
    return ws


def _validation(ws, header: str, headers: List[str], options: List[str]):
    if header not in headers or ws.max_row < 2:
        return
    col = get_column_letter(headers.index(header) + 1)
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
    dv.error, dv.errorTitle = "Listeden seçin", "Geçersiz değer"
    ws.add_data_validation(dv)
    dv.add("%s2:%s%d" % (col, col, max(ws.max_row, 2) + 500))


def export_workbook(summary: List[list], screens: List[list], tables: List[list], columns: List[list],
                    relations: List[list], inventory: List[list]) -> bytes:
    wb = Workbook()
    ws = _sheet(wb, "Özet", ["Bilgi", "Değer"], summary, {1: 34, 2: 60}, first=True)
    ws.freeze_panes = None

    _sheet(wb, "Ekranlar", ["Ekran", "Tür", "Modül", "Başlangıç", "Tablo sayısı", "Kolon sayısı", "SQL sayısı", "Not"],
           screens, {1: 40, 3: 24, 8: 40})

    t_headers = ["Şema", "Tablo", "Tip", "Satır", "DB açıklaması", "Siron modülü", "Araştırmaya göre", "Araştırma güveni",
                 "Ekranlar", "Kullanılan kolon", "Durum", "Hedef", "Filtre", "Not"]
    ws = _sheet(wb, "Tablolar", t_headers, tables, {2: 30, 5: 36, 6: 30, 7: 44, 8: 16, 9: 40, 12: 28, 13: 32, 14: 48},
                editable=["Durum", "Hedef", "Filtre", "Not"])
    _validation(ws, "Durum", t_headers, list(STATUS_LABELS.values()))

    c_headers = ["Şema", "Tablo", "Kolon", "Tip", "DB açıklaması", "Kullanım", "Ekranlar", "Doluluk %",
                 "Farklı değer", "Örnek değerler", "Durum", "Hassas", "Hedef alan", "Not"]
    ws = _sheet(wb, "Kolon eşleme", c_headers, columns,
                {2: 26, 3: 26, 5: 36, 7: 36, 10: 36, 13: 30, 14: 48},
                editable=["Durum", "Hassas", "Hedef alan", "Not"])
    _validation(ws, "Durum", c_headers, list(STATUS_LABELS.values()))
    _validation(ws, "Hassas", c_headers, ["Evet", "Hayır"])

    _sheet(wb, "İlişkiler", ["Alt tablo", "Alt kolonlar", "Üst tablo", "Üst kolonlar", "Güven", "Kaynak",
                             "Veri kontrolü"], relations, {1: 30, 2: 24, 3: 30, 4: 24, 6: 30})
    i_headers = ["ID", "Modül", "Ekran", "Açıklama", "Öncelik", "Durum", "Yakalama sayısı", "Not"]
    _sheet(wb, "Envanter", i_headers, inventory, {2: 24, 3: 40, 4: 40, 8: 40},
           editable=["Modül", "Ekran", "Açıklama", "Öncelik", "Durum", "Not"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _rows(ws) -> List[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h or "").strip() for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def _text(v) -> Optional[str]:
    if v is None:
        return None
    text = str(v).strip()
    return text or None


def parse_status(v) -> Optional[str]:
    text = _text(v)
    if not text:
        return None
    code = STATUS_CODES.get(text.casefold())
    if code is None:
        raise ValueError("Bilinmeyen durum: %s" % text)
    return code


def parse_flag(v) -> Optional[int]:
    text = _text(v)
    if not text or "tahmin" in text.casefold():  # "Evet (tahmin)" = automatic detection, not a decision
        return None
    if text.casefold() in YES:
        return 1
    if text.casefold() in NO:
        return 0
    raise ValueError("Hassas için Evet/Hayır bekleniyor: %s" % text)


def read_workbook(data: bytes) -> dict:
    """{'tables': [...], 'columns': [...], 'inventory': [...], 'errors': [...]} of normalised rows."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = {"tables": [], "columns": [], "inventory": [], "errors": []}
    if "Tablolar" in wb.sheetnames:
        for i, r in enumerate(_rows(wb["Tablolar"]), start=2):
            try:
                out["tables"].append({"owner": _text(r.get("Şema")), "table": _text(r.get("Tablo")),
                                      "status": parse_status(r.get("Durum")), "target": _text(r.get("Hedef")),
                                      "filter": _text(r.get("Filtre")), "note": _text(r.get("Not"))})
            except ValueError as exc:
                out["errors"].append("Tablolar satır %d: %s" % (i, exc))
    if "Kolon eşleme" in wb.sheetnames:
        for i, r in enumerate(_rows(wb["Kolon eşleme"]), start=2):
            try:
                out["columns"].append({"owner": _text(r.get("Şema")), "table": _text(r.get("Tablo")),
                                       "column": _text(r.get("Kolon")), "status": parse_status(r.get("Durum")),
                                       "sensitive": parse_flag(r.get("Hassas")), "target": _text(r.get("Hedef alan")),
                                       "note": _text(r.get("Not"))})
            except ValueError as exc:
                out["errors"].append("Kolon eşleme satır %d: %s" % (i, exc))
    if "Envanter" in wb.sheetnames:
        for r in _rows(wb["Envanter"]):
            if _text(r.get("Ekran")):
                out["inventory"].append({"id": r.get("ID"), "module": _text(r.get("Modül")),
                                         "name": _text(r.get("Ekran")), "description": _text(r.get("Açıklama")),
                                         "priority": _text(r.get("Öncelik")), "status": _text(r.get("Durum")),
                                         "notes": _text(r.get("Not"))})
    return out
