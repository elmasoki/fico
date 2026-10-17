import io

import openpyxl

from siron_tracer.sqlutil import bind_segments, build_slots, find_placeholders, parse_audit_binds, substitute_binds
from siron_tracer.sqlutil import compute_sql_id
from siron_tracer.timeline import build_steps
from test_app import Q1, Q2, rec
from test_features import capture, env  # noqa: F401  (pytest fixture)

S1, S2 = compute_sql_id(Q1), compute_sql_id(Q2)


def test_bind_segments_match_substitution():
    sql = "SELECT * FROM t WHERE a = :1 AND b = :name AND c = :name AND d = :3"
    ph = find_placeholders(sql)
    slots = build_slots(ph, parse_audit_binds(" #1(4):1001 #2(5):AHMET"))
    segs = bind_segments(sql, slots)
    assert "".join(t for t, _ in segs) == substitute_binds(sql, ph, slots)
    assert ("'1001'", 0) in segs and ("'AHMET'", 1) in segs
    assert (":3", 3) in segs and slots[3].missing                      # missing value keeps the placeholder
    formatted = bind_segments(sql, slots, formatted=True)
    assert "\n" in "".join(t for t, _ in formatted) and ("'1001'", 0) in formatted
    assert bind_segments("SELECT 1 FROM dual", []) == [("SELECT 1 FROM dual", None)]


def test_timeline_steps():
    def r(i, sec, sql, binds=None):
        return {"id": i, "ts": "2026-10-17 10:00:%06.3f" % sec + "000", "sql_id": compute_sql_id(sql), "sql_text": sql,
                "sql_binds": binds, "action": "SELECT"}
    records = [r(1, 0.0, Q1, " #1(1):1"), r(2, 0.4, Q1, " #1(1):1"), r(3, 0.9, Q2, " #1(1):A #2(1):B"),
               r(4, 6.0, Q1, " #1(1):2"), r(5, 6.2, "SELECT 1 FROM DUAL")]
    steps = build_steps(records, {S1: [("SIRON", "CUSTOMER")], S2: [("SIRON", "ALERT")]}, gap=2.0,
                        noise_ids={compute_sql_id("SELECT 1 FROM DUAL")})
    assert len(steps) == 2
    assert steps[0].items[0].repeat == 2 and steps[0].sql_count == 3       # same SQL + binds folded
    assert steps[0].tables == [("SIRON", "CUSTOMER"), ("SIRON", "ALERT")]
    assert steps[1].gap_before == 5.1 and steps[1].start == 6.0 and steps[1].items[-1].noise
    assert steps[1].items[0].binds == ["2"]
    hidden = build_steps(records, {}, gap=2.0, noise_ids={compute_sql_id("SELECT 1 FROM DUAL")}, hide_noise=True)
    assert all(not i.noise for s in hidden for i in s.items)
    assert len(build_steps(records, {}, gap=10)) == 1


def test_link_sql_to_screen_and_inventory(env):
    client, oracle, store, _ = env
    client.post("/inventory", data={"action": "add", "name": "Müşteri Detay", "module": "Müşteri"})
    sid = store.list_screens()[0]["id"]
    cid = capture(client, store, oracle, [rec(1, Q1, " #1(1):7"), rec(2, Q2, " #1(1):A #2(1):B")], name="Detay açılışı")
    page = client.get("/captures/%d/sql/%s" % (cid, S1)).get_data(as_text=True)
    assert "Ekran envanterine bağla" in page

    client.post("/inventory/sql", data={"capture_id": cid, "sql_id": S1, "screen_id": sid, "note": "Hesaplar sekmesi"})
    link = store.screen_sqls(screen_id=sid)[0]
    assert link["sql_id"] == S1 and link["capture_name"] == "Detay açılışı" and link["note"] == "Hesaplar sekmesi"
    assert link["sql_text"] == Q1 and link["capture_start"]
    assert store.get_capture(cid)["screen_id"] == sid                       # unlinked capture adopts the screen
    client.post("/inventory/sql", data={"capture_id": cid, "sql_id": S1, "screen_id": sid, "note": "Hesaplar"})
    assert len(store.screen_sqls()) == 1 and store.screen_sqls()[0]["note"] == "Hesaplar"   # relink updates note

    html = client.get("/inventory").get_data(as_text=True)
    c = store.get_capture(cid)
    assert "Detay açılışı" in html and c["start_local"][:16] in html and S1 in html and "Hesaplar" in html
    html = client.get("/inventory/%d" % sid).get_data(as_text=True)
    assert S1 in html and "Detay açılışı" in html and "SIRON.CUSTOMER" in html and "bağlı SQL" in html
    assert client.get("/inventory/9999").status_code == 404
    assert "Müşteri Detay" in client.get("/captures/%d/sql/%s" % (cid, S1)).get_data(as_text=True)
    assert 'selected' in client.get("/captures?screen=%d" % sid).get_data(as_text=True)

    client.post("/inventory/sql", data={"capture_id": cid, "sql_id": S2, "new_screen": "Alarm > Alarm Listesi"})
    alarm = next(s for s in store.list_screens() if s["name"] == "Alarm Listesi")
    assert alarm["module"] == "Alarm" and store.screen_sqls(screen_id=alarm["id"])[0]["sql_id"] == S2

    wb = openpyxl.load_workbook(io.BytesIO(client.get("/mapping.xlsx").data))
    rows = list(wb["Envanter"].values)
    assert rows[0][-1] == "Bağlı SQL'ler" and any(S1 in (r[-1] or "") for r in rows[1:])

    client.post("/inventory/sql/%d/delete" % link["id"])
    assert store.screen_sqls(screen_id=sid) == []
    client.post("/inventory", data={"action": "delete", "id": alarm["id"]})
    assert store.screen_sqls() == []
    client.post("/inventory/sql", data={"capture_id": cid, "sql_id": S1, "screen_id": sid})
    client.post("/captures/%d/delete" % cid)
    assert store.screen_sqls() == []                                        # links go with the capture


def test_bind_values_highlighted_and_manual_edit_kept(env):
    client, oracle, store, _ = env
    cid = capture(client, store, oracle, [rec(1, Q2, " #1(4):OPEN")])
    html = client.get("/captures/%d/sql/%s" % (cid, S2)).get_data(as_text=True)
    assert '<mark class="bv" data-slot="0"' in html and "&#39;OPEN&#39;<sup>1</sup>" in html
    assert 'class="bv missing" data-slot="1"' in html                      # :B2 has no audited value
    edited = "SELECT * FROM alert WHERE status = 'CLOSED'"
    html = client.post("/captures/%d/sql/%s" % (cid, S2),
                       data={"do": "run", "sql_override": edited, "literal_orig": "x", "value_0": "OPEN",
                             "kind_0": "text", "value_1": "", "kind_1": "text"}).get_data(as_text=True)
    assert oracle.ran[-1][0] == edited and "elle düzenlendi" in html
    client.post("/captures/%d/sql/%s" % (cid, S2),                          # second run keeps the manual edit
                data={"do": "run", "sql_override": edited, "literal_orig": "", "value_0": "OPEN", "kind_0": "text",
                      "value_1": "", "kind_1": "text"})
    assert oracle.ran[-1][0] == edited


def test_timeline_page_and_bulk_link(env):
    client, oracle, store, _ = env
    cid = capture(client, store, oracle, [rec(1, Q1, " #1(1):1"), rec(9, Q2, " #1(1):A #2(1):B")])
    html = client.get("/captures/%d/timeline" % cid).get_data(as_text=True)
    assert "Adım 1" in html and "Adım 2" in html and S1 in html and "Zaman akışı" in html
    assert "Adım 2" not in client.get("/captures/%d/timeline?gap=30" % cid).get_data(as_text=True)
    client.post("/inventory/sql/bulk", data={"capture_id": cid, "sql": [S1, S2, S1], "new_screen": "Müşteri > Detay",
                                             "note": "Açılış"})
    links = store.screen_sqls()
    assert sorted(l["sql_id"] for l in links) == sorted([S1, S2]) and links[0]["note"] == "Açılış"
    assert "bağlı" in client.get("/captures/%d/timeline" % cid).get_data(as_text=True)
    r = client.post("/inventory/sql/bulk", data={"capture_id": cid})
    assert r.status_code == 302 and len(store.screen_sqls()) == 2


def test_bind_value_search(env):
    client, oracle, store, _ = env
    cid = capture(client, store, oracle, [rec(1, Q1, " #1(6):123456"), rec(2, Q2, " #1(4):OPEN #2(5):C0042")])
    html = client.get("/search?q=c0042").get_data(as_text=True)
    assert "Bind değerlerinde" in html and S2 in html and ">C0042</mark>" in html
    assert S1 not in client.get("/search?q=C0042").get_data(as_text=True).split("Bind değerlerinde")[1].split("SQL'ler")[0]
    html = client.get("/search?q=6").get_data(as_text=True)                 # "(6)" header digits do not count
    hits = html.split("Bind değerlerinde")[1].split("</table>")[0]
    assert S1 in hits and S2 not in hits
    page = client.post("/captures/%d/find" % cid, data={"value": "123456", "mode": "exact"}).get_data(as_text=True)
    assert "parametre olarak gitmiş" in page and S1 in page
    page = client.post("/captures/%d/find" % cid, data={"value": "2345", "mode": "exact"}).get_data(as_text=True)
    assert "parametre olarak gitmiş" not in page
    page = client.post("/captures/%d/find" % cid, data={"value": "2345", "mode": "contains"}).get_data(as_text=True)
    assert "parametre olarak gitmiş" in page


def test_compare_captures(env):
    client, oracle, store, _ = env
    a = capture(client, store, oracle, [rec(1, Q1, " #1(1):1")], name="Detay öncesi")
    b = capture(client, store, oracle, [rec(2, Q1, " #1(1):2"), rec(3, Q2, " #1(1):A #2(1):B")], name="Detay sonrası")
    html = client.get("/captures/compare?c=%d&c=%d" % (a, b)).get_data(as_text=True)
    only_b = html.split("Yalnız B")[1]
    assert "SIRON.ALERT" in html.split("yalnız B'de</p>")[1].split("</div>")[0]
    assert S2 in only_b and "Detay öncesi" in html and "Detay sonrası" in html
    assert "SIRON.CUSTOMER" in html.split("iki yakalamada da</p>")[1].split("</div>")[0]
    assert "SIRON.ALERT" not in html.split("yalnız A'da</p>")[1].split("</div>")[0]
    assert client.get("/captures/compare").status_code == 200
    assert "Karşılaştır" in client.get("/captures").get_data(as_text=True)
