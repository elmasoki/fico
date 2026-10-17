import io
import zipfile

import pytest

from siron_tracer import create_app
from siron_tracer.config import Settings
from siron_tracer.excel import read_workbook
from siron_tracer.sqlutil import compute_sql_id
from siron_tracer.store import Store
from test_app import CATALOG, FakeOracle, Q1, Q2, rec


@pytest.fixture
def env(tmp_path):
    settings = Settings(oracle_dsn="x", oracle_user="u", siron_db_user="SIRON", siron_schemas=["SIRON"],
                        secret_key="t", export_dir=str(tmp_path / "exports"))
    oracle = FakeOracle()
    app = create_app(settings, oracle=oracle, store=Store(":memory:"), sync_jobs=True)
    app.config["TESTING"] = True
    client = app.test_client()
    client.post("/catalog/refresh")
    return client, oracle, app.extensions["siron"]["store"], app


def capture(client, store, oracle, records, **form):
    oracle.records = records
    data = {"name": "Ekran"}
    data.update(form)
    client.post("/captures", data=data)
    cid = store.list_captures()[0]["id"]
    client.post("/captures/%d/stop" % cid)
    return cid


def test_multiple_users_and_schema(env):
    client, oracle, store, _ = env
    cid = capture(client, store, oracle, [rec(1, Q1, " #1(1):1")], db_users=["SIRON"],
                  new_users="rapor_user, BATCH_USER", new_schema="siron_kyc", kind="batch")
    c = store.get_capture(cid)
    assert c["db_users"] == ["SIRON", "RAPOR_USER", "BATCH_USER"]
    assert c["schema_name"] == "SIRON_KYC" and c["kind"] == "batch"
    assert oracle.last_filter.db_users == ["SIRON", "RAPOR_USER", "BATCH_USER"]
    assert store.get_list("db_users") == ["RAPOR_USER", "BATCH_USER"]
    assert store.get_list("schemas") == ["SIRON_KYC"]
    html = client.get("/captures").get_data(as_text=True)
    assert "RAPOR_USER" in html and "SIRON_KYC" in html            # offered in the form next time
    client.post("/captures/%d/sql/%s" % (cid, compute_sql_id(Q1)), data={"do": "run"})
    assert oracle.last_schema == "SIRON_KYC"                         # queries run in the capture schema
    client.post("/captures/%d/update" % cid, data={"name": "X", "kind": "screen", "schema": "SIRON"})
    assert store.get_capture(cid)["schema_name"] == "SIRON"


def test_inventory_and_coverage(env):
    client, oracle, store, _ = env
    client.post("/inventory", data={"action": "bulk", "lines": "Müşteri > Müşteri Detay\nAlarm > Alarm Listesi\nAlarm\tVaka"})
    screens = store.list_screens()
    assert {(s["module"], s["name"]) for s in screens} == {("Müşteri", "Müşteri Detay"), ("Alarm", "Alarm Listesi"), ("Alarm", "Vaka")}
    sid = next(s["id"] for s in screens if s["name"] == "Müşteri Detay")
    capture(client, store, oracle, [rec(1, Q1, " #1(1):1")], screen_id=str(sid), name="")
    assert store.list_captures()[0]["name"] == "Müşteri Detay"
    html = client.get("/inventory").get_data(as_text=True)
    assert "Yakalandı" in html
    html = client.get("/coverage").get_data(as_text=True)
    assert "33%" in html.replace(",", ".") or "33" in html          # 1 of 3 screens
    assert "SIRON.ALERT" in html                                     # never touched by a capture
    assert client.get("/").status_code == 200
    capture(client, store, oracle, [rec(1, Q2, " #1(1):A #2(1):B")], name="Yeni ekran", add_to_inventory="1")
    assert any(s["name"] == "Yeni ekran" for s in store.list_screens())


def test_jobs_profile_and_dependencies(env):
    client, oracle, store, _ = env
    assert store.list_jobs(1)[0]["kind"] == "catalog" and store.list_jobs(1)[0]["status"] == "done"
    r = client.post("/table/SIRON/CUSTOMER", data={"action": "profile"})
    job_id = int(r.headers["Location"].rsplit("/", 1)[1])
    assert store.get_job(job_id)["status"] == "done"
    assert client.get("/jobs/%d" % job_id).status_code == 200
    assert client.get("/jobs/%d.json" % job_id).get_json()["status"] == "done"
    assert client.get("/jobs/active.json").get_json() == []
    assert store.profiles(("SIRON", "CUSTOMER"))[("SIRON", "CUSTOMER", "NAME")]["fill_pct"] == 90.0
    html = client.get("/table/SIRON/CUSTOMER").get_data(as_text=True)
    assert "90,0%" in html
    # a SQL calling a package reaches the package's tables indirectly
    cid = capture(client, store, oracle, [rec(1, "SELECT pkg_risk.score(:1) FROM dual", " #1(1):1")])
    html = client.get("/captures/%d" % cid).get_data(as_text=True)
    assert "SIRON.ALERT" in html and "dolaylı" in html
    assert "PKG_RISK" in client.get("/table/SIRON/ALERT").get_data(as_text=True)
    client.post("/tables/bulk", data={"k": ["SIRON.ALERT", "SIRON.ACCOUNT"], "action": "profile"})
    assert store.profiles(("SIRON", "ALERT"))


def test_literal_variants_are_grouped(env):
    client, oracle, store, _ = env
    cid = capture(client, store, oracle, [rec(1, "SELECT name FROM customer WHERE id = 1"),
                                          rec(2, "SELECT name FROM customer WHERE id = 2"),
                                          rec(3, Q1, " #1(1):1")])
    html = client.get("/captures/%d" % cid).get_data(as_text=True)
    assert "+1 varyant" in html
    assert "+1 varyant" not in client.get("/captures/%d?group=0" % cid).get_data(as_text=True)
    sid = compute_sql_id("SELECT name FROM customer WHERE id = 1")
    assert "literal varyantı" in client.get("/captures/%d/sql/%s" % (cid, sid)).get_data(as_text=True)


def test_excel_round_trip_and_undo(env):
    client, oracle, store, _ = env
    capture(client, store, oracle, [rec(1, Q1, " #1(1):1")])
    client.post("/table/SIRON/CUSTOMER", data={"action": "annotate", "status": "review"})
    data = client.get("/mapping.xlsx").data
    parsed = read_workbook(data)
    assert any(t["table"] == "CUSTOMER" and t["status"] == "review" for t in parsed["tables"])
    assert any(c["column"] == "IBAN" for c in parsed["columns"])

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Kolon eşleme"]
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[headers.index("Kolon")].value == "NAME":
            row[headers.index("Durum")].value = "Taşınacak"
            row[headers.index("Hedef alan")].value = "PARTY.FULL_NAME"
            row[headers.index("Hassas")].value = "Evet"
    wst = wb["Tablolar"]
    th = [c.value for c in wst[1]]
    for row in wst.iter_rows(min_row=2):
        if row[th.index("Tablo")].value == "CUSTOMER":
            row[th.index("Durum")].value = "Taşınacak"
            row[th.index("Filtre")].value = "ID > 0"
    out = io.BytesIO()
    wb.save(out)
    client.post("/excel/import", data={"file": (io.BytesIO(out.getvalue()), "esleme.xlsx")},
                content_type="multipart/form-data")
    anns = store.annotations()
    assert anns[("SIRON", "CUSTOMER", "NAME")]["target"] == "PARTY.FULL_NAME"
    assert anns[("SIRON", "CUSTOMER", "NAME")]["sensitive"] == 1
    assert anns[("SIRON", "CUSTOMER", "")]["status"] == "migrate"
    assert anns[("SIRON", "CUSTOMER", "")]["extract_filter"] == "ID > 0"
    assert "esleme.xlsx" in client.get("/excel").get_data(as_text=True)
    client.post("/excel/undo")
    anns = store.annotations()
    assert ("SIRON", "CUSTOMER", "NAME") not in anns and anns[("SIRON", "CUSTOMER", "")]["status"] == "review"

    # invalid status values are rejected as a whole
    ws["K2"].value = "Belki"
    bad = io.BytesIO()
    wb.save(bad)
    r = client.post("/excel/import", data={"file": (io.BytesIO(bad.getvalue()), "bad.xlsx")},
                    content_type="multipart/form-data", follow_redirects=True)
    assert "hata var" in r.get_data(as_text=True)


def test_migration_package_and_extract(env):
    client, oracle, store, _ = env
    capture(client, store, oracle, [rec(1, Q1, " #1(1):1"), rec(2, Q2, " #1(1):A #2(1):B")])
    client.post("/tables/bulk", data={"k": ["SIRON.ACCOUNT", "SIRON.CUSTOMER"], "action": "migrate"})
    client.post("/table/SIRON/ACCOUNT", data={"action": "annotate", "status": "migrate", "extract_filter": "ID < 100",
                                              "status_IBAN": "ignore"})
    html = client.get("/migration").get_data(as_text=True)
    assert html.index("SIRON.CUSTOMER") < html.index("SIRON.ACCOUNT")   # parent first
    z = zipfile.ZipFile(io.BytesIO(client.get("/migration.zip").data))
    names = set(z.namelist())
    assert {"manifest.csv", "extract_sqlplus.sql", "counts.sql", "kolon_esleme.csv", "README.txt",
            "sql/01_CUSTOMER.sql", "sql/02_ACCOUNT.sql"} <= names
    client.post("/table/SIRON/CUSTOMER", data={"action": "annotate", "status": "migrate", "status_NAME": "migrate"})
    z = zipfile.ZipFile(io.BytesIO(client.get("/migration.zip").data))
    customer_sql = z.read("sql/01_CUSTOMER.sql").decode()
    assert "NAME" in customer_sql and "ID" in customer_sql      # primary key travels even if not marked
    assert "CUST_NO" not in customer_sql
    account_sql = z.read("sql/02_ACCOUNT.sql").decode()
    assert "CUST_ID" in account_sql                              # relation column to the parent
    assert "WHERE ID < 100" in account_sql and "IBAN" not in account_sql
    r = client.post("/migration/extract", data={"mode": "decided"})
    job = store.get_job(int(r.headers["Location"].rsplit("/", 1)[1]))
    assert job["status"] == "done" and len(job["result"]["files"]) == 2
    f = job["result"]["files"][0]
    body = client.get("/exports/%s/%s" % (job["result"]["run"], f["file"])).get_data(as_text=True)
    assert "Ayşe" in body
    assert client.get("/exports/../x/y").status_code == 404


def test_activity_log_settings_and_check(env):
    client, oracle, store, _ = env
    capture(client, store, oracle, [rec(1, Q1, " #1(1):1")])
    client.post("/sql", data={"sql": "select 1 from dual", "do": "run"})
    client.get("/mapping.csv")
    actions = {a["action"] for a in store.activity()}
    assert {"catalog", "query", "export_csv"} <= actions
    assert "select 1 from dual" in client.get("/activity?format=csv").get_data(as_text=True)
    assert "Sorgu çalıştırma" in client.get("/activity").get_data(as_text=True)
    client.post("/settings", data={"db_users": "a_user\nb_user", "schemas": "S2"})
    assert store.get_list("db_users") == ["A_USER", "B_USER"] and store.get_list("schemas") == ["S2"]
    assert "A_USER" in client.get("/settings").get_data(as_text=True)


def test_delete_captures(env):
    client, oracle, store, _ = env
    a = capture(client, store, oracle, [rec(1, Q1, " #1(1):1")], name="Ekran A")
    b = capture(client, store, oracle, [rec(2, Q2, None)], name="Ekran B")
    c = capture(client, store, oracle, [rec(3, Q1, " #1(1):2")], name="Ekran C")
    store.save_annotation("SIRON", "CUSTOMERS", "", status="migrate")
    html = client.get("/captures").get_data(as_text=True)
    assert "Seçilenleri sil" in html and "/captures/%d/delete" % a in html

    client.post("/captures/%d/delete" % a)                          # single, from the list row
    assert store.get_capture(a) is None and store.records(a) == []
    client.post("/captures/delete", data={"c": [str(b), str(c), "9999"]})   # bulk; unknown ids ignored
    assert store.list_captures() == [] and store.records(b) == []
    assert store.annotations()[("SIRON", "CUSTOMERS", "")]["status"] == "migrate"   # decisions survive
    deletes = [x["detail"] for x in store.activity() if x["action"] == "capture_delete"]
    assert any("Ekran A" in d for d in deletes) and any("Ekran B" in d and "Ekran C" in d for d in deletes)
    assert "Yakalama silme" in client.get("/activity").get_data(as_text=True)
    resp = client.post("/captures/delete", data={})
    assert resp.status_code == 302 and store.list_captures() == []


def test_reset_running_capture(env):
    client, oracle, store, _ = env
    oracle.records = [rec(1, Q1, " #1(1):1")]
    client.post("/captures", data={"name": "Müşteri detay"})
    cid = store.list_captures()[0]["id"]
    first = store.get_capture(cid)
    page = client.get("/captures/%d" % cid).get_data(as_text=True)
    assert "/captures/%d/reset" % cid in page and "Sıfırla" in page

    client.post("/captures/%d/reset" % cid)
    after = store.get_capture(cid)
    assert after["status"] == "recording"                           # still recording
    assert after["start_utc"] > first["start_utc"] and after["start_local"] > first["start_local"]
    live = client.get("/captures/%d/live" % cid).get_json()
    assert oracle.last_filter.start_utc.strftime("%Y-%m-%d %H:%M:%S.%f") == after["start_utc"]
    assert live["elapsed"] == 60                                    # FakeOracle clock ticks a minute per call
    assert any(a["action"] == "capture_reset" for a in store.activity())

    client.post("/captures/%d/stop" % cid)                          # stop fetches only the new window
    assert oracle.last_filter.start_utc.strftime("%Y-%m-%d %H:%M:%S.%f") == after["start_utc"]
    assert store.get_capture(cid)["status"] == "done"
    client.post("/captures/%d/reset" % cid)                         # finished captures cannot be reset
    assert store.get_capture(cid)["start_utc"] == after["start_utc"]
    assert store.get_capture(cid)["record_count"] == 1


def test_csrf_protection(env):
    client, oracle, store, app = env
    app.config["TESTING"] = False
    assert client.post("/settings", data={"db_users": "x"}).status_code == 400
    page = client.get("/settings").get_data(as_text=True)
    import re
    token = re.search(r'var token = "([^"]+)"', page).group(1)
    assert client.post("/settings", data={"db_users": "x", "_csrf": token}).status_code == 302


def test_every_page_renders(env):
    client, oracle, store, _ = env
    cid = capture(client, store, oracle, [rec(1, Q1, " #1(1):1"), rec(2, Q2, " #1(1):A #2(1):B")])
    client.post("/inventory", data={"action": "add", "name": "Ekran 1", "module": "M"})
    sid = compute_sql_id(Q1)
    for url in ["/", "/captures", "/captures?kind=batch", "/captures/%d" % cid, "/captures/%d/columns" % cid,
                "/captures/%d/find" % cid, "/captures/%d/sql/%s" % (cid, sid), "/captures/%d/sql/%s?cursor=1" % (cid, sid),
                "/inventory", "/coverage", "/tables", "/tables?all=1", "/table/SIRON/CUSTOMER", "/table/-/NOPE",
                "/mapping", "/mapping?status=none&sensitive=1", "/relations", "/relations?state=unverified",
                "/relations?capture=%d" % cid, "/er", "/er?capture=%d" % cid, "/builder", "/sql", "/migration",
                "/report", "/report.html", "/report.html?capture=%d&cols=all" % cid, "/excel", "/snapshots",
                "/jobs", "/activity", "/settings", "/check", "/search?q=cust", "/search"]:
        r = client.get(url)
        assert r.status_code == 200, url
        assert "Traceback" not in r.get_data(as_text=True)
